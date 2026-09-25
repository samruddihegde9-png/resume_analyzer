"""
Document parsing: turns an uploaded PDF/DOCX/TXT file into clean text plus
structured contact info. No AI models here — pure extraction + regex.
"""

import io
import re
from dataclasses import dataclass, field
from typing import List, Optional

import pdfplumber
import docx  # python-docx

from src.utils import clean_text, normalize_whitespace

SECTION_HEADINGS = [
    "SUMMARY", "OBJECTIVE", "PROFILE",
    "EDUCATION", "ACADEMIC BACKGROUND",
    "EXPERIENCE", "WORK EXPERIENCE", "PROFESSIONAL EXPERIENCE", "EMPLOYMENT HISTORY",
    "INTERNSHIPS", "INTERNSHIP",
    "PROJECTS", "PERSONAL PROJECTS", "ACADEMIC PROJECTS",
    "SKILLS", "TECHNICAL SKILLS", "CORE COMPETENCIES",
    "CERTIFICATIONS", "CERTIFICATES", "LICENSES",
    "ACHIEVEMENTS", "AWARDS", "HONORS",
    "PUBLICATIONS",
    "EXTRACURRICULAR", "ACTIVITIES", "VOLUNTEER", "VOLUNTEERING",
    "LANGUAGES",
    "CONTACT",
]

# canonical section name -> heading variants that map to it
SECTION_CANONICAL = {
    "Summary": ["SUMMARY", "OBJECTIVE", "PROFILE"],
    "Education": ["EDUCATION", "ACADEMIC BACKGROUND"],
    "Experience": ["EXPERIENCE", "WORK EXPERIENCE", "PROFESSIONAL EXPERIENCE", "EMPLOYMENT HISTORY"],
    "Internships": ["INTERNSHIPS", "INTERNSHIP"],
    "Projects": ["PROJECTS", "PERSONAL PROJECTS", "ACADEMIC PROJECTS"],
    "Skills": ["SKILLS", "TECHNICAL SKILLS", "CORE COMPETENCIES"],
    "Certifications": ["CERTIFICATIONS", "CERTIFICATES", "LICENSES"],
    "Achievements": ["ACHIEVEMENTS", "AWARDS", "HONORS"],
    "Publications": ["PUBLICATIONS"],
    "Extracurricular": ["EXTRACURRICULAR", "ACTIVITIES", "VOLUNTEER", "VOLUNTEERING"],
    "Languages": ["LANGUAGES"],
}

IMPORTANT_SECTIONS = ["Summary", "Education", "Experience", "Skills", "Projects"]


@dataclass
class ContactInfo:
    name: Optional[str] = None
    email: Optional[str] = None
    phone: Optional[str] = None
    linkedin: Optional[str] = None
    github: Optional[str] = None
    portfolio: Optional[str] = None


@dataclass
class ParsedDocument:
    raw_text: str
    cleaned_text: str
    lines: List[str] = field(default_factory=list)
    contact: ContactInfo = field(default_factory=ContactInfo)
    sections: dict = field(default_factory=dict)  # canonical name -> text
    section_order: List[str] = field(default_factory=list)
    extraction_quality: str = "good"  # good / partial / poor
    warnings: List[str] = field(default_factory=list)


# ------------------------------------------------------------------
# Raw text extraction per file type
# ------------------------------------------------------------------

def _extract_pdf_text(file_bytes: bytes) -> str:
    text = ""
    try:
        import fitz  # PyMuPDF — imported lazily so the app still runs
                     # (with a reduced PDF path via pdfplumber below) if
                     # PyMuPDF isn't installed in a given environment.
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        for page in doc:
            text += page.get_text()
        doc.close()
    except Exception:
        text = ""

    if not text.strip():
        try:
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                for page in pdf.pages:
                    extracted = page.extract_text()
                    if extracted:
                        text += extracted + "\n"
        except Exception:
            pass
    return text


def _extract_docx_text(file_bytes: bytes) -> str:
    text_parts = []
    try:
        document = docx.Document(io.BytesIO(file_bytes))
        for para in document.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    text_parts.append(" | ".join(cells))
    except Exception:
        pass
    return "\n".join(text_parts)


def _extract_txt_text(file_bytes: bytes) -> str:
    for encoding in ("utf-8", "latin-1"):
        try:
            return file_bytes.decode(encoding)
        except Exception:
            continue
    return ""


def extract_text(file_bytes: bytes, filename: str) -> (str, str):
    """Returns (raw_text, extraction_quality)."""
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""

    if ext == "pdf":
        raw = _extract_pdf_text(file_bytes)
    elif ext in ("docx", "doc"):
        raw = _extract_docx_text(file_bytes)
    elif ext == "txt":
        raw = _extract_txt_text(file_bytes)
    else:
        raw = ""

    word_count = len(raw.split())
    if word_count == 0:
        quality = "poor"
    elif word_count < 60:
        quality = "partial"
    else:
        quality = "good"

    return raw, quality


# ------------------------------------------------------------------
# Contact info extraction
# ------------------------------------------------------------------

EMAIL_RE = re.compile(r"[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}")
PHONE_RE = re.compile(r"(\+?\d{1,3}[\s\-.]?)?(\(?\d{2,4}\)?[\s\-.]?)?\d{3,4}[\s\-.]?\d{3,4}")
LINKEDIN_RE = re.compile(r"(https?://)?(www\.)?linkedin\.com/[a-zA-Z0-9\-/_%]+", re.IGNORECASE)
GITHUB_RE = re.compile(r"(https?://)?(www\.)?github\.com/[a-zA-Z0-9\-/_%]+", re.IGNORECASE)
PORTFOLIO_RE = re.compile(
    r"(https?://)?(www\.)?[a-zA-Z0-9\-]+\.(dev|me|io|com|xyz|tech|site|vercel\.app|netlify\.app|github\.io)\b[^\s]*",
    re.IGNORECASE,
)


def _guess_name(lines: List[str]) -> Optional[str]:
    for line in lines[:8]:
        candidate = line.strip()
        if not candidate or len(candidate) > 60:
            continue
        if EMAIL_RE.search(candidate) or any(ch.isdigit() for ch in candidate):
            continue
        words = candidate.split()
        if 1 < len(words) <= 5 and all(w[0].isupper() or not w[0].isalpha() for w in words if w):
            return normalize_whitespace(candidate)
    return None


def extract_contact_info(raw_text: str, lines: List[str]) -> ContactInfo:
    email_match = EMAIL_RE.search(raw_text)
    linkedin_match = LINKEDIN_RE.search(raw_text)
    github_match = GITHUB_RE.search(raw_text)

    portfolio = None
    for match in PORTFOLIO_RE.finditer(raw_text):
        url = match.group()
        if "linkedin.com" in url.lower() or "github.com" in url.lower():
            continue
        portfolio = url
        break

    phone = None
    for match in PHONE_RE.finditer(raw_text):
        digits = re.sub(r"\D", "", match.group())
        if 7 <= len(digits) <= 13:
            phone = match.group().strip()
            break

    return ContactInfo(
        name=_guess_name(lines),
        email=email_match.group() if email_match else None,
        phone=phone,
        linkedin=linkedin_match.group() if linkedin_match else None,
        github=github_match.group() if github_match else None,
        portfolio=portfolio,
    )


# ------------------------------------------------------------------
# Section splitting
# ------------------------------------------------------------------

def _heading_to_canonical(line: str) -> Optional[str]:
    stripped = re.sub(r"[^A-Za-z& ]", "", line).strip().upper()
    if not stripped or len(stripped) > 40:
        return None
    for canonical, variants in SECTION_CANONICAL.items():
        for variant in variants:
            if stripped == variant or stripped.startswith(variant):
                return canonical
    return None


def split_sections(raw_text: str) -> (dict, List[str]):
    """Split resume text into canonical sections by scanning for heading-like
    lines (short, upper-case-ish, matches known heading vocabulary)."""
    lines = raw_text.split("\n")
    sections = {}
    order = []
    current = "Header"
    buffer = []

    def flush():
        if buffer:
            sections.setdefault(current, "")
            sections[current] += "\n".join(buffer) + "\n"

    for line in lines:
        candidate = _heading_to_canonical(line)
        is_heading_like = candidate is not None and (
            line.strip().isupper() or len(line.strip().split()) <= 4
        )
        if is_heading_like:
            flush()
            current = candidate
            if current not in order:
                order.append(current)
            buffer = []
        else:
            buffer.append(line)
    flush()

    for k in sections:
        sections[k] = sections[k].strip()

    return sections, order


# ------------------------------------------------------------------
# Public entry point
# ------------------------------------------------------------------

def parse_resume(file_bytes: bytes, filename: str) -> ParsedDocument:
    raw_text, quality = extract_text(file_bytes, filename)
    cleaned = clean_text(raw_text)
    lines = [l for l in raw_text.split("\n") if l.strip()]

    warnings = []
    if quality == "poor":
        warnings.append(
            "Very little text could be extracted. The file may be a scanned "
            "image or have an unusual layout."
        )
    elif quality == "partial":
        warnings.append(
            "Limited text was extracted. Results may be incomplete — "
            "double-check formatting if this looks short."
        )

    contact = extract_contact_info(raw_text, lines)
    sections, order = split_sections(raw_text)

    return ParsedDocument(
        raw_text=raw_text,
        cleaned_text=cleaned,
        lines=lines,
        contact=contact,
        sections=sections,
        section_order=order,
        extraction_quality=quality,
        warnings=warnings,
    )


def parse_plain_text(text: str) -> ParsedDocument:
    """Used for job descriptions pasted directly as text."""
    cleaned = clean_text(text)
    lines = [l for l in text.split("\n") if l.strip()]
    sections, order = split_sections(text)
    return ParsedDocument(
        raw_text=text,
        cleaned_text=cleaned,
        lines=lines,
        contact=ContactInfo(),
        sections=sections,
        section_order=order,
        extraction_quality="good" if len(text.split()) > 20 else "partial",
        warnings=[],
    )
